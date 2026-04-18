import os
from pathlib import Path

from docx import Document


path = Path(os.environ["DOCX_PATH"])
doc = Document(path)
table = doc.tables[0]
table.cell(15, 0).text = "候选池"
table.cell(15, 1).text = "候选方案规模"
table.cell(15, 2).text = "N_c"
table.cell(15, 3).text = "2417/1149"
table.cell(15, 4).text = "个"
table.cell(15, 5).text = "工程效率/资源回收模式候选总数"
table.cell(15, 6).text = "planningResults"
doc.save(path)
print(f"patched table1 {path}")
